"""Growth Lab PowerPoint builder — editable .pptx on the official GL template.

Rendering layer of the `gl-pptx` skill. Decks are built on the official team template
(`assets/GL_presentation_template.potx`) whose **12 layouts are the ground truth for
PowerPoint**: logos are baked into every layout and all placeholder text inherits
Source Sans Pro from the slide master.

What this module is:
  - a builder per **slide class** from `recipes/slide.md`, each mapped onto one of the
    template's 12 layouts (see `references/slide-classes.md` for the mapping);
  - the grammar's tokens (`grammar.md` section 1) and the slide recipe's type scale
    (`recipes/slide.md` section 2) as Python constants, converted px -> pt;
  - `copy_slide()` — deep-copy a slide between decks preserving images and LIVE charts
    (used by the compiler);
  - the `figs/` pipeline (`export_fig` / `find_fig` / `render_figure`) — charts arrive as
    images, produced in any language;
  - `validate_deck()` — re-open a finished deck and check it is valid, on-template and
    on-grammar.

Two deliberate per-medium compromises, documented here because this file is a downstream
encoding of grammar.md (see the README's drift table):

  1. **Type stack.** The grammar's stack is Source Serif 4 (voice) + Inter (function).
     This medium keeps the template master's **Source Sans Pro** instead: the 12 layouts
     are the PowerPoint ground truth, logo and placeholder styling are baked into them,
     and a deck must render on any machine that opens it. Builders therefore **never set
     a font** — placeholder text inherits the master. The grammar's serif/sans split is
     carried by **weight and italic** instead: headings semibold, chart source italic,
     eyebrows uppercase + tracked.
  2. **Optical sizing.** `opsz` is not expressible in OOXML, so the grammar's optical-size
     axis is unavailable here (grammar section 2 says such stubs default to text optical
     size).

Everything else — the ink ramp, accent, the 6x3-tone categorical palette, the sequential
and diverging ramps, the dark-tone-for-text rule, table rules, and the type SCALE — follows
grammar.md and recipes/slide.md exactly. The px -> pt conversion is lossless: the template
canvas is 13.333 x 7.5 in = 960 x 540 pt, which is the recipe's 1280 x 720 px at 96 DPI,
so every recipe value is exactly x0.75.

Never hardcode a hex in a deck script — import the tokens from here.
"""
from __future__ import annotations

import copy
import io
import json
import os
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

# ───────────────────────── paths ─────────────────────────
_HERE = Path(__file__).resolve()
SKILL_DIR = _HERE.parents[1]                       # .../skills/gl-pptx/
TEMPLATE_PATH = SKILL_DIR / "assets" / "GL_presentation_template.potx"


def design_root() -> Path | None:
    """Locate the gl-design plugin root (holds grammar.md + recipes/).

    Resolution order mirrors `theme_gl.R`: GL_DESIGN_ROOT -> CLAUDE_PLUGIN_ROOT -> this
    script's own location. Returns None when the skill has been copied out of the plugin
    (the builders still work; only the drift check needs the root).
    """
    for cand in (os.environ.get("GL_DESIGN_ROOT"),
                 os.environ.get("CLAUDE_PLUGIN_ROOT"),
                 _HERE.parents[3]):
        if cand and (Path(cand) / "grammar.md").exists():
            return Path(cand)
    return None


def default_output_dir() -> Path:
    """Where decks are written when no path is given: $GL_SLIDES_DIR, else the cwd."""
    env = os.environ.get("GL_SLIDES_DIR")
    return Path(env) if env else Path.cwd()


# ───────────────────────── grammar tokens (grammar.md section 1) ─────────────────────────
# A downstream copy, per the repo's "values flow downward" rule. `check_token_drift()`
# verifies every hex below still appears in grammar.md.
INK = {"ink": "#1A1714", "ink-2": "#2C2823", "ink-3": "#4F4A42", "ink-4": "#9A9389"}
ACCENT = {"accent": "#1A5A8E", "accent-deep": "#003E6B",
          "accent-soft": "#3A85B8", "accent-tint": "#E1F0FA"}
PAPER = {"paper": "#FFFFFF", "cover-bg": "#F3F2EA", "cover-disk": "#ECEBE0",
         "paper-warm": "#F4F1EA", "rule": "#DDDDDD", "gridline": "#D8D4CC"}

# Categorical palette — six hues, three tones each. Main fills; dark strokes AND all text
# associated with the mark; light for backgrounds/faded states (grammar section 3.3).
CAT = {
    "c-1": {"light": "#B5D5EA", "main": "#2F87C8", "dark": "#1A5A8E"},
    "c-2": {"light": "#E89C9C", "main": "#CC4948", "dark": "#8A2C2B"},
    "c-3": {"light": "#92D6BF", "main": "#2AA584", "dark": "#1A6B53"},
    "c-4": {"light": "#B5A0CC", "main": "#7554A3", "dark": "#4A3470"},
    "c-5": {"light": "#F4BC8A", "main": "#EA822D", "dark": "#A8580F"},
    "c-6": {"light": "#E6E2A8", "main": "#CDC86B", "dark": "#8A8638"},
}
MUTED = {"light": "#CDD2D9", "main": "#AFB5BE", "dark": "#5F6773"}

SEQUENTIAL = {
    "sequential-1": ["#E5F0F9", "#B5D5EA", "#6FA5CE", "#2F87C8", "#1A5A8E"],
    "sequential-2": ["#F4D5D5", "#E89C9C", "#DC6F6E", "#CC4948", "#8A2C2B"],
    "sequential-3": ["#D5EFE7", "#92D6BF", "#5BC0A0", "#2AA584", "#1A6B53"],
    "sequential-4": ["#E5DDF0", "#B5A0CC", "#9276BA", "#7554A3", "#4A3470"],
    "sequential-5": ["#FBE5D5", "#F4BC8A", "#EE9A52", "#EA822D", "#A8580F"],
    "sequential-6": ["#FBF8DC", "#E6E2A8", "#DCD68E", "#CDC86B", "#8A8638"],
}
DIVERGING = {
    "div-2-1": ["#8A2C2B", "#DC6F6E", "#EFC7C0", "#C5DCEC", "#6FA5CE", "#1A5A8E"],
    "div-3-1": ["#1A6B53", "#5BC0A0", "#BDE5D8", "#C5DCEC", "#6FA5CE", "#1A5A8E"],
    "div-5-1": ["#A8580F", "#EE9A52", "#F4BC8A", "#C5DCEC", "#6FA5CE", "#1A5A8E"],
    "div-6-1": ["#8A8638", "#DCD68E", "#E6E2A8", "#C5DCEC", "#6FA5CE", "#1A5A8E"],
}

# Semantic aliases (grammar section 3.1) — the institutional voice and the lead-finding red.
HIGHLIGHT = CAT["c-1"]["main"]
LEAD_FINDING = CAT["c-2"]["main"]


def token(name: str) -> str:
    """Resolve a grammar token name to its hex: 'ink-2', 'accent', 'c-3.dark',
    'c-muted', 'c-muted-dark', 'paper', 'gridline', 'sequential-1.2'."""
    if name in INK:
        return INK[name]
    if name in ACCENT:
        return ACCENT[name]
    if name in PAPER:
        return PAPER[name]
    if name == "c-muted":
        return MUTED["main"]
    if name in ("c-muted-light", "c-muted-dark"):
        return MUTED[name.rsplit("-", 1)[1]]
    if "." in name:
        base, part = name.rsplit(".", 1)
        if base in CAT:
            return CAT[base][part]
        if base in SEQUENTIAL:
            return SEQUENTIAL[base][int(part)]
        if base in DIVERGING:
            return DIVERGING[base][int(part)]
    if name in CAT:
        return CAT[name]["main"]
    raise KeyError("unknown grammar token: " + repr(name))


def categorical(n: int, tone: str = "main") -> list:
    """First `n` categorical colors in grammar order (c-1 first). Six is a hard ceiling
    (grammar rule 5) — past that, mute the rest instead of adding hues."""
    if n > 6:
        raise ValueError(
            str(n) + " categorical colors requested; grammar rule 5 caps a chart at 6. "
            "Mute the rest with MUTED and let one or two focus series carry the story.")
    return [CAT["c-" + str(i)][tone] for i in range(1, n + 1)]


def emphasis(categories, focus, focus_token: str = "c-1") -> dict:
    """Highlight-by-muting (grammar section 3.1 / rule 4): focus series in c-1 (or c-2 for
    a lead finding), everything else in c-muted."""
    focus_set = {focus} if isinstance(focus, str) else set(focus)
    hi = CAT[focus_token]["main"]
    return {c: (hi if c in focus_set else MUTED["main"]) for c in categories}


def label_color(fill_hex: str) -> str:
    """The dark tone that must be used for any text or stroke naming a mark filled with
    `fill_hex` (grammar rule 6 — no exceptions, including the muted grey)."""
    up = fill_hex.upper()
    for tones in list(CAT.values()) + [MUTED]:
        if up in (tones["main"].upper(), tones["light"].upper(), tones["dark"].upper()):
            return tones["dark"]
    return INK["ink-2"]


def check_token_drift(root: Path | None = None) -> dict:
    """Verify every hex in this file still appears in grammar.md (README, "Auditing for
    drift"). A mismatch is a bug HERE, never in grammar.md."""
    root = root or design_root()
    out = {"ok": True, "checked": 0, "missing": [], "grammar": None}
    if root is None:
        out["ok"] = False
        out["missing"].append("grammar.md not found - set GL_DESIGN_ROOT")
        return out
    text = (root / "grammar.md").read_text(encoding="utf-8").upper()
    out["grammar"] = str(root / "grammar.md")
    hexes = {}
    for gname, g in (("ink", INK), ("accent", ACCENT), ("paper", PAPER), ("muted", MUTED)):
        for k, v in g.items():
            hexes[gname + "." + k] = v
    for k, tones in CAT.items():
        for t, v in tones.items():
            hexes[k + "." + t] = v
    for name, ramp in dict(**SEQUENTIAL, **DIVERGING).items():
        for i, v in enumerate(ramp):
            hexes[name + "[" + str(i) + "]"] = v
    for name, v in hexes.items():
        out["checked"] += 1
        if v.upper() not in text:
            out["missing"].append(name + " = " + v)
    out["ok"] = not out["missing"]
    return out


# ───────────────── type scale (recipes/slide.md section 2, px x 0.75) ─────────────────
# The template canvas is 960 x 540 pt = the recipe's 1280 x 720 px at 96 DPI, so the
# conversion is exact. Layout placeholders inherit their size from the template; these
# apply to the text this module DRAWS (manual textboxes, tables, captions).
def px(v: float) -> float:
    """Recipe px -> pptx pt (exact, x0.75)."""
    return v * 0.75


PT = {
    "cover_title": px(64), "break_title": px(56), "closing_title": px(40),
    "h1": px(40), "h2": px(26), "h3": px(20),
    "body": px(24), "blockquote": px(22),
    "chart_title": px(26), "chart_source": px(16),
    "eyebrow": px(13), "figure_label": px(13),
    "table_header": px(20), "table_cell": px(20),
    "folio": px(11), "cover_date": px(14), "closing_tag": px(13),
}

# Recipe section 1 padding per slide class, in pt (px x 0.75). Applies to the areas this
# module draws; the template's own placeholders keep their baked-in geometry.
PAD = {
    "content":  (px(64), px(80), px(72), px(80)),
    "chart":    (px(48), px(64), px(56), px(64)),
    "break":    (px(80), px(96), px(80), px(96)),
    "img_full": (px(40), px(56), px(40), px(56)),
    "map":      (px(28), px(40), px(28), px(40)),
    "closing":  (px(80), px(96), px(80), px(96)),
}
TRACKING = {"eyebrow": 0.14, "cover_date": 0.18, "closing_tag": 0.18}   # em (recipe section 2)

# ───────────────────────── canvas ─────────────────────────
SLIDE_W = Emu(12192000)                 # 13.333 in
SLIDE_H = Emu(6858000)                  # 7.5 in
MARGIN = Pt(px(80))                     # recipe content padding, left/right
CONTENT_W = SLIDE_W - 2 * MARGIN
BODY_TOP = Inches(1.3)                  # just below the template's title placeholder
FOOTER_SAFE_BOTTOM = Inches(6.82)       # the template's footer band starts here
TITLE_ZONE_TOP = Inches(0.40)


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str.lstrip("#").upper())


# ─────────────── layout names — the 12, exactly as the template spells them ───────────────
L_TITLE = "Title Slide"
L_TITLE_BG = "Title Slide With Background Image"
L_SIDE = "Side Image + Text"
L_SINGLE = "Single Visual"
L_TWO = "Two Visuals "                  # NOTE trailing space in the template
L_VISUAL_TEXT = "1_Two Visuals "        # one visual + a text column
L_STATEMENT = "Statement Text"
L_FULL = "Full Image "                  # NOTE trailing space
L_TITLE_BLANK = "Title + Blank"
L_BLANK = "Blank"
L_BLANK_NF = "Blank without footer"
L_CLOSING = "Closing Slide"

TEMPLATE_LAYOUTS = [L_TITLE, L_TITLE_BG, L_SIDE, L_SINGLE, L_TWO, L_VISUAL_TEXT,
                    L_STATEMENT, L_FULL, L_TITLE_BLANK, L_BLANK, L_BLANK_NF, L_CLOSING]

# recipe slide class -> (builder, template layout). See references/slide-classes.md.
SLIDE_CLASSES = {
    "title":     ("add_title_slide", L_TITLE + " / " + L_TITLE_BG),
    "content":   ("add_content_slide", L_TITLE_BLANK),
    "chart":     ("add_chart_slide", L_SINGLE),
    "cols":      ("add_cols_slide", L_TWO + " / " + L_VISUAL_TEXT),
    "map":       ("add_map_slide", L_TWO),
    "img_slide": ("add_img_slide", L_SIDE),
    "img_full":  ("add_img_full_slide", L_FULL),
    "break":     ("add_break_slide", L_STATEMENT),
    "table":     ("add_table_slide", L_TITLE_BLANK),
    "closing":   ("add_closing_slide", L_CLOSING),
    "blank":     ("add_blank_slide", L_BLANK + " / " + L_BLANK_NF),
}

# ───────────────────────── template loading ─────────────────────────
_TEMPLATE_CT = b"presentationml.template.main+xml"
_PRES_CT = b"presentationml.presentation.main+xml"


def _load_template(template: Path | None = None) -> Presentation:
    """Open the GL template, rewriting the .potx content-type to a presentation in memory
    so python-pptx accepts it. Works for .potx and .pptx alike."""
    template = Path(template) if template else TEMPLATE_PATH
    if not template.exists():
        raise FileNotFoundError(
            "GL template not found at " + str(template) + ". gl-pptx needs "
            "GL_presentation_template.potx in its assets/ folder.")
    raw = template.read_bytes()
    buf = io.BytesIO()
    seen = set()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin, \
            zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in seen:          # some .potx zips carry duplicate entries
                continue
            seen.add(item.filename)
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_TEMPLATE_CT, _PRES_CT)
            zout.writestr(item, data)
    buf.seek(0)
    return Presentation(buf)


def _clear_slides(prs: Presentation):
    """Drop the template's instructional slide(s): remove each slide's relationship (and
    thus its orphaned part) as well as its entry in the id list, so the leftover
    slide1.xml part cannot collide with newly added slides."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
        sldIdLst.remove(sldId)


def new_deck(template: Path | None = None) -> Presentation:
    """An empty 16:9 deck on the GL template (instructional slides stripped)."""
    prs = _load_template(template)
    if not (prs.slide_width and abs(int(prs.slide_width) - int(SLIDE_W)) < 50000):
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _clear_slides(prs)
    return prs


def layout_names(prs: Presentation) -> list:
    return [lay.name for lay in prs.slide_layouts]


def _layout(prs: Presentation, name: str):
    """Find a layout by name, tolerant of the template's trailing spaces and case."""
    want = name.strip().lower()
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    for lay in prs.slide_layouts:
        if lay.name.strip().lower() == want:
            return lay
    raise KeyError("layout " + repr(name) + " not in template; have " + str(layout_names(prs)))


def _add(prs: Presentation, layout_name: str):
    return prs.slides.add_slide(_layout(prs, layout_name))


# ───────────────────────── placeholders ─────────────────────────
_PIC_TYPES = {PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.BITMAP, PP_PLACEHOLDER.OBJECT}


def _ph(slide, idx: int):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def _ph_or_kind(slide, idx: int, kind: str = "text", nth: int = 0):
    """Placeholder by index, falling back to the nth placeholder of a kind. The indices
    used by the builders come from the shipped template; this fallback keeps them working
    if a future template revision renumbers them."""
    ph = _ph(slide, idx)
    if ph is not None:
        return ph
    pool = []
    for p in slide.placeholders:
        is_pic = p.placeholder_format.type in _PIC_TYPES
        if (kind == "picture") == bool(is_pic):
            pool.append(p)
    return pool[nth] if len(pool) > nth else None


def _set_ph(slide, idx: int, text, nth: int = 0):
    """Fill a placeholder's text. The run inherits the layout/master font — we never set
    a font here (see the module docstring, compromise 1)."""
    ph = _ph_or_kind(slide, idx, "text", nth)
    if ph is None:
        return None
    tf = ph.text_frame
    items = [text] if isinstance(text, str) else list(text)
    for i, item in enumerate(items):
        t, level = (item if isinstance(item, tuple) else (item, 0))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        try:
            p.level = level
        except Exception:
            pass
        p.text = t
    return ph


def _drop(shape):
    shape._element.getparent().remove(shape._element)


def _place_image(slide, left, top, width, height, image_path, mode="contain"):
    """Add a picture fitted into a rect. mode='contain' (whole image, never cropped — for
    charts, per the recipe's chart slide) or 'cover' (fill, cropping overflow — photos)."""
    iw = ih = None
    try:
        from PIL import Image
        with Image.open(image_path) as im:
            iw, ih = im.size
    except Exception:
        pass
    L, T, W, H = int(left), int(top), int(width), int(height)
    if not (iw and ih):                                    # no Pillow: best effort
        return slide.shapes.add_picture(
            str(image_path), Emu(L), Emu(T),
            width=Emu(W), height=Emu(H) if mode == "cover" else None)
    ra, ia = W / H, iw / ih
    if mode == "cover":
        pic = slide.shapes.add_picture(str(image_path), Emu(L), Emu(T),
                                       width=Emu(W), height=Emu(H))
        if ia > ra:
            pic.crop_left = pic.crop_right = (1 - ra / ia) / 2
        elif ia < ra:
            pic.crop_top = pic.crop_bottom = (1 - ia / ra) / 2
        return pic
    scale = min(W / iw, H / ih)
    dw, dh = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(str(image_path), Emu(L + (W - dw) // 2),
                                    Emu(T + (H - dh) // 2), width=Emu(dw), height=Emu(dh))


def _fill_image_ph(slide, idx: int, image_path, mode="contain", nth: int = 0,
                   top_offset: int = 0):
    """Place an image into a picture placeholder's rect, then remove the now-empty
    placeholder so its prompt text cannot show.

    `top_offset` (EMU) shifts the top down and shortens the rect by the same amount, so
    the bottom edge stays put — used when an eyebrow has grown the title block. The rect is
    read BEFORE the placeholder is dropped and never written back: assigning `.top` to a
    placeholder that inherits its position from the layout writes a partial `<a:xfrm>` and
    the geometry comes back wrong.
    """
    ph = _ph_or_kind(slide, idx, "picture", nth)
    if ph is None:
        return None
    left, top, width, height = int(ph.left), int(ph.top), int(ph.width), int(ph.height)
    _drop(ph)
    if top_offset:
        top += top_offset
        height = max(int(Inches(1.0)), height - top_offset)
    return _place_image(slide, left, top, width, height, image_path, mode=mode)


# ───────────────────────── text we draw ourselves ─────────────────────────
def _textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return box, tf


def _style_run(run, size: float, color: str = None, bold=False, italic=False,
               caps=False, tracking: float = 0.0):
    """Style a run we created. `size` in pt (from PT), `color` a grammar token hex.

    No font name is ever set — the master's Source Sans Pro is inherited. Weight, italic,
    caps and tracking carry the grammar's role distinctions in this medium.
    """
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = _rgb(color or INK["ink-2"])
    rPr = run._r.get_or_add_rPr()
    if caps:
        rPr.set("cap", "all")
    if tracking:
        rPr.set("spc", str(int(round(tracking * size * 100))))   # em -> 1/100 pt


def _para(tf, first: bool, text: str, size: float, align=None, space_after: float = 0,
          **kw):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after:
        p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _style_run(run, size, **kw)
    return p


def _eyebrow(tf, first: bool, text: str):
    """Eyebrow / figure label: Inter 600 accent UPPERCASE, 0.14em tracking in the grammar;
    here uppercase + bold + tracked in the master font."""
    return _para(tf, first, text, PT["eyebrow"], color=ACCENT["accent"], bold=True,
                 caps=True, tracking=TRACKING["eyebrow"], space_after=px(6))


EYEBROW_H = Pt(px(24))          # extra title-box height an eyebrow line needs


def _set_title_with_eyebrow(slide, idx: int, title: str, eyebrow: str = "") -> int:
    """Fill a title placeholder, optionally with an eyebrow line above the title INSIDE
    the placeholder. Returns the EMU the box grew by (0 when there is no eyebrow).

    The eyebrow cannot be its own textbox above the title: the template's title
    placeholder starts at 0.27 in, so there is no room above it. Putting the eyebrow in
    the placeholder as a preceding paragraph gives the recipe's compact eyebrow-over-title
    block and keeps the title inheriting the master's styling — but the box must also grow
    downward, or the title is pushed onto whatever sits below it.
    """
    ph = _ph_or_kind(slide, idx, "text")
    if ph is None:
        return 0
    tf = ph.text_frame
    if not eyebrow:
        tf.paragraphs[0].text = title
        return 0
    _eyebrow(tf, True, eyebrow)
    p = tf.add_paragraph()
    p.text = title                       # inherits the placeholder's title styling
    delta = int(EYEBROW_H)
    # Write the FULL rect (left/top/width/height), not just the height: a placeholder that
    # inherits its position from the layout has no <a:xfrm> of its own, and setting one
    # dimension writes a partial one — which collapses the box to a sliver.
    left, top, width, height = int(ph.left), int(ph.top), int(ph.width), int(ph.height)
    ph.left, ph.top, ph.width, ph.height = left, top, width, height + delta
    return delta


def _source_line(slide, text: str, top=None):
    """Chart source: Source Serif 4 italic in the grammar; italic ink-2 here. Sits above
    the template's footer band."""
    top = top if top is not None else Inches(6.42)
    _box, tf = _textbox(slide, MARGIN, top, CONTENT_W, Pt(px(28)))
    _para(tf, True, text, PT["chart_source"], color=INK["ink-2"], italic=True)
    return _box


_BULLET_CHARS = {0: "•", 1: "–", 2: "–"}


def _set_bullet(paragraph, level: int, size: float):
    """Give a paragraph a real PowerPoint bullet with a hanging indent.

    Prepending "• " to the text instead (the obvious shortcut) leaves wrapped lines
    starting underneath the glyph, so a two-line bullet loses its left edge. `marL` with a
    negative `indent` is what makes continuation lines align with the text.
    """
    pPr = paragraph._p.get_or_add_pPr()
    step = int(Pt(px(26)))                        # one indent step per level
    hang = int(Pt(px(22)))                        # glyph-to-text distance
    pPr.set("marL", str(step * (level + 1)))
    pPr.set("indent", str(-hang))
    # OOXML child order: spcBef/spcAft come first, then buFont, then buChar.
    for tag, attrs in (("a:buFont", {"typeface": "Arial", "panose": "020B0604020202020204",
                                     "pitchFamily": "34", "charset": "0"}),
                       ("a:buChar", {"char": _BULLET_CHARS.get(level, "–")})):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
        pPr.append(pPr.makeelement(qn(tag), attrs))


def _body_textbox(slide, left, top, width, height, body, size=None, bullets=True):
    """A body block: 24px/18pt, ink-2, one item per paragraph. `body` is a string or a
    list of strings / (text, level) tuples."""
    size = size or PT["body"]
    _box, tf = _textbox(slide, left, top, width, height)
    items = [body] if isinstance(body, str) else list(body)
    for i, item in enumerate(items):
        t, level = (item if isinstance(item, tuple) else (item, 0))
        p = _para(tf, i == 0, t, size, color=INK["ink-2"],
                  space_after=px(9))       # recipe section 3: list item 0.35em
        if bullets and len(items) > 1:
            _set_bullet(p, level, size)
    return _box


# ───────────────────── slide builders — one per recipe slide class ─────────────────────
def add_title_slide(prs, title: str, subtitle: str = "", footer: str = "",
                    background=None):
    """`title` -> **Title Slide** (or **Title Slide With Background Image** when
    `background=` is given). The layout owns the cover's artwork, logo and type — we only
    fill it. Recipe rule: cover meta is the date alone; pass it as `subtitle`."""
    s = _add(prs, L_TITLE_BG if background else L_TITLE)
    _set_ph(s, 0, title)
    if subtitle:
        _set_ph(s, 10, subtitle, nth=1)
    if footer:
        _set_ph(s, 11, footer, nth=2)
    if background:
        _fill_image_ph(s, 12, background, mode="cover")
    return s


def add_content_slide(prs, title: str, body, source: str = ""):
    """`content` -> **Title + Blank** + a body block. One idea per slide (recipe rule 2)."""
    s = _add(prs, L_TITLE_BLANK)
    _set_ph(s, 14, title)
    _body_textbox(s, MARGIN, BODY_TOP, CONTENT_W, Inches(4.9), body)
    if source:
        _source_line(s, source)
    return s


def add_chart_slide(prs, title: str, image_path, source: str = "", eyebrow: str = ""):
    """`chart` -> **Single Visual**. Compact eyebrow + finding-style title at the top, the
    chart contained (never cropped), source line below. Chart titles end with a period
    (recipe rule 5) — `validate_deck` flags ones that do not.

    The slide owns the title and source, so a deck-bound chart image should carry NEITHER
    baked in (that is what gl-matplotlib's `chart_text` is for, on standalone PNGs).
    """
    s = _add(prs, L_SINGLE)
    delta = _set_title_with_eyebrow(s, 14, title, eyebrow)
    # an eyebrow grows the title block, so the chart starts lower and keeps its bottom
    _fill_image_ph(s, 16, image_path, mode="contain", top_offset=delta)
    if source:
        if _ph(s, 15) is not None:
            _set_ph(s, 15, source)
        else:
            _source_line(s, source)
    return s


def add_cols_slide(prs, title: str, left, right, source: str = ""):
    """`cols` -> **Two Visuals** (two images) or **1_Two Visuals** (image + text column).

    `left`/`right` may each be an image path or text (str/list). Use for parallel
    comparisons, not to force two unrelated ideas onto one slide (recipe section 4).
    """
    left_is_img = _looks_like_image(left)
    right_is_img = _looks_like_image(right)
    if left_is_img and right_is_img:
        s = _add(prs, L_TWO)
        _set_ph(s, 14, title)
        mode = "contain"
        _fill_image_ph(s, 17, left, mode=mode, nth=0)
        _fill_image_ph(s, 18, right, mode=mode, nth=1)
    else:
        s = _add(prs, L_VISUAL_TEXT)
        _set_ph(s, 14, title)
        img, txt = (left, right) if left_is_img else (right, left)
        _fill_image_ph(s, 17, img, mode="contain")
        _set_ph(s, 18, txt if isinstance(txt, str) else [t for t in txt], nth=1)
    if source:
        if _ph(s, 15) is not None:
            _set_ph(s, 15, source)
        else:
            _source_line(s, source)
    return s


def add_map_slide(prs, title: str, left, right, source: str = ""):
    """`map` -> **Two Visuals**, but with the recipe's tight map padding (28/40 px instead
    of the content 64/80): side-by-side maps want every pixel, so the two images are placed
    into computed rects rather than the layout's picture placeholders."""
    s = _add(prs, L_TWO)
    _set_ph(s, 14, title)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type in _PIC_TYPES:
            _drop(ph)
    pad_t, pad_r, pad_b, pad_l = PAD["map"]
    top = Inches(1.15)
    gutter = Pt(px(40))
    bottom = FOOTER_SAFE_BOTTOM - Pt(pad_b)
    height = int(bottom) - int(top)
    width = (int(SLIDE_W) - int(Pt(pad_l)) - int(Pt(pad_r)) - int(gutter)) // 2
    _place_image(s, Pt(pad_l), top, width, height, left, mode="contain")
    _place_image(s, int(Pt(pad_l)) + width + int(gutter), top, width, height, right,
                 mode="contain")
    if source:
        _source_line(s, source)
    return s


def add_img_slide(prs, heading: str, body, image_path):
    """`img_slide` -> **Side Image + Text** — text/bullets beside a side image
    (cover-cropped)."""
    s = _add(prs, L_SIDE)
    _set_ph(s, 15, heading)
    _set_ph(s, 16, body, nth=1)
    _fill_image_ph(s, 13, image_path, mode="cover")
    return s


def add_img_full_slide(prs, image_path, caption: str = ""):
    """`img_full` -> **Full Image** — full-bleed image, optional caption line below."""
    s = _add(prs, L_FULL)
    _fill_image_ph(s, 18, image_path, mode="cover")
    if caption:
        _source_line(s, caption)
    return s


def add_break_slide(prs, text: str):
    """`break` -> **Statement Text** — a section divider / single big takeaway.

    Recipe rule 7: no body and no chart on a break slide; it is a punctuation mark.
    """
    s = _add(prs, L_STATEMENT)
    _set_ph(s, 14, text)
    return s


def add_table_slide(prs, title: str, dataframe, source: str = "", max_rows: int = 12):
    """`table` -> **Title + Blank** + a grammar-styled table.

    Grammar tables: ink header (sentence case, never caps), ink-2 cells, top/bottom ink
    rules with hairline `rule` dividers, no fills, numeric columns right-aligned with
    tabular figures. The recipe fits ~6 rows x 4 cols comfortably — past that, split.
    """
    s = _add(prs, L_TITLE_BLANK)
    _set_ph(s, 14, title)
    df = dataframe.head(max_rows)
    n_rows, n_cols = df.shape[0] + 1, df.shape[1]
    row_h = Pt(px(38))
    gfx = s.shapes.add_table(n_rows, n_cols, MARGIN, BODY_TOP, CONTENT_W,
                             Emu(int(row_h) * n_rows))
    table = gfx.table
    _strip_table_style(table)
    table.first_row = False
    table.horz_banding = False
    numeric = [_is_numeric_col(df, c) for c in df.columns]
    # Precision is decided per COLUMN, from its largest value: mixing 31.5 with 5.17 in one
    # column reads as sloppiness, and a column is only scannable if its decimal points line
    # up.
    fmts = [_column_fmt(df[c]) if numeric[j] else None
            for j, c in enumerate(df.columns)]
    for j, col in enumerate(df.columns):
        _set_cell_text(table.cell(0, j), str(col), PT["table_header"], INK["ink"],
                       bold=True, right=numeric[j])
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for j, val in enumerate(row):
            text = fmts[j](val) if fmts[j] else _fmt(val)
            _set_cell_text(table.cell(i, j), text, PT["table_cell"], INK["ink-2"],
                           right=numeric[j])
    _table_rules(table)
    if source:
        _source_line(s, source)
    return s


def add_closing_slide(prs, message: str = "", contact: str = ""):
    """`closing` -> **Closing Slide** — the branded closer.

    This layout is already complete: it has NO placeholders, a solid dark ground, and the
    lab/HKS lockup with the URL as baked-in shapes. So the default is to use it as it comes.

    A `message` / `contact` is overlaid only when asked for, and then in `paper` — the
    grammar's ink tones are unreadable on this ground — in the upper area, clear of the
    centred lockup (which occupies roughly 3.0-4.0 in).
    """
    s = _add(prs, L_CLOSING)
    if message or contact:
        _box, tf = _textbox(s, MARGIN, Inches(0.95), CONTENT_W, Inches(1.5),
                            anchor=MSO_ANCHOR.TOP)
        first = True
        if message:
            _para(tf, True, message, PT["closing_title"], color=PAPER["paper"],
                  italic=True, align=PP_ALIGN.CENTER, space_after=px(18))
            first = False
        if contact:
            _para(tf, first, contact, PT["closing_tag"], color=PAPER["paper"], caps=True,
                  tracking=TRACKING["closing_tag"], align=PP_ALIGN.CENTER)
    return s


def add_blank_slide(prs, footer: bool = True):
    """`blank` -> **Blank** / **Blank without footer** — freeform canvas, and the
    copy-as-is target for the compiler."""
    return _add(prs, L_BLANK if footer else L_BLANK_NF)


def _looks_like_image(v) -> bool:
    if isinstance(v, (str, Path)):
        return str(v).lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif",
                                       ".tiff", ".emf", ".wmf"))
    return False


# ───────────────────────── table internals ─────────────────────────
def _column_fmt(series):
    """A single formatter for one numeric column, so every cell in it shows the same
    number of decimals — chosen from the column's largest magnitude."""
    try:
        import math
        vals = [abs(float(v)) for v in series
                if v is not None and not (isinstance(v, float) and math.isnan(v))]
        top = max(vals) if vals else 0.0
        all_int = all(float(v).is_integer() for v in series
                      if v is not None and not (isinstance(v, float) and math.isnan(v)))
    except Exception:
        return _fmt
    dec = 0 if (all_int or top >= 1000) else (1 if top >= 10 else 2)

    def f(val):
        try:
            import math
            if val is None or (isinstance(val, float) and math.isnan(val)):
                return ""
            return "{:,.{}f}".format(float(val), dec)
        except Exception:
            return _fmt(val)
    return f


def _is_numeric_col(df, col) -> bool:
    try:
        import pandas as pd
        return bool(pd.api.types.is_numeric_dtype(df[col]))
    except Exception:
        return False


def _set_cell_text(cell, text, size, color, bold=False, right=False):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Pt(px(10))
    cell.margin_right = Pt(px(10))
    cell.margin_top = Pt(px(4))
    cell.margin_bottom = Pt(px(4))
    cell.fill.background()                      # grammar tables carry no fills
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if right:
        p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    _style_run(run, size, color=color, bold=bold)


def _strip_table_style(table):
    """Drop the inherited PowerPoint table style so the grammar's rules are what show
    (no banding, no accent-filled header)."""
    tblPr = table._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            tblPr.remove(style)


def _cell_borders(cell, left=None, right=None, top=None, bottom=None):
    """Set all four edges of a cell in one pass (python-pptx has no border API).

    Each edge is either None -> explicitly NO line, or (hex, width_pt) -> a solid line.
    Both matter: once the table style id is stripped, PowerPoint falls back to its own
    default look and draws a full grid unless every unwanted edge is explicitly `noFill`.

    The four elements MUST appear in the order lnL, lnR, lnT, lnB, ahead of the rest of
    tcPr — OOXML validates child order, and out-of-order edges make PowerPoint discard the
    whole spec (which is exactly how the vertical rules survived the first attempt). They
    are inserted at index 0 in reverse so the final order is correct.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for tag, spec in (("a:lnB", bottom), ("a:lnT", top),
                      ("a:lnR", right), ("a:lnL", left)):
        if spec is None:
            ln = tcPr.makeelement(qn(tag), {"w": "0"})
            ln.append(ln.makeelement(qn("a:noFill"), {}))
        else:
            hex_color, width_pt = spec
            ln = tcPr.makeelement(qn(tag), {"w": str(int(width_pt * 12700)),
                                            "cap": "flat", "cmpd": "sng", "algn": "ctr"})
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(ln.makeelement(qn("a:srgbClr"),
                                       {"val": hex_color.lstrip("#").upper()}))
            ln.append(fill)
        tcPr.insert(0, ln)


def _table_rules(table):
    """Grammar table rules: horizontal only — 1pt ink above and below the header, 1pt ink
    under the last row, hairline `rule` dividers between body rows, and NO vertical rules
    or outer box (columns are separated by space, not lines)."""
    n_rows = len(table.rows)
    n_cols = len(table.columns)
    ink = (INK["ink"], 1.0)
    hair = (PAPER["rule"], 0.5)
    for j in range(n_cols):
        for i in range(n_rows):
            if i == 0:
                top, bottom = ink, ink                       # header, both rules
            elif i == n_rows - 1:
                top, bottom = None, ink                      # last row closes the table
            else:
                top, bottom = None, hair                     # hairline divider
            _cell_borders(table.cell(i, j), left=None, right=None, top=top, bottom=bottom)


def _fmt(val):
    """Format a cell value for a slide: thousands separators, and precision scaled to
    magnitude rather than a fixed two decimals.

    A slide is read across a room, so `33,841.96` spends four glyphs on precision nobody
    can use and makes the column harder to scan. Significant digits, not decimal places:
    thousands round to whole units, values under ten keep two decimals.
    """
    try:
        import math
        if val is None:
            return ""
        if isinstance(val, bool):
            return str(val)
        if isinstance(val, float) and math.isnan(val):
            return ""
        if isinstance(val, int) or (isinstance(val, float) and val == int(val)):
            return "{:,}".format(int(val))
        if isinstance(val, float):
            a = abs(val)
            if a >= 1000:
                return "{:,.0f}".format(val)
            if a >= 10:
                return "{:,.1f}".format(val)
            return "{:,.2f}".format(val)
    except Exception:
        pass
    return "" if val is None else str(val)


# ───────────────────── cross-deck slide copy (the compiler) ─────────────────────
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def copy_slide(src_slide, dest_prs, layout=None):
    """Deep-copy a slide from another presentation into `dest_prs`, preserving shapes,
    embedded images and LIVE charts. `layout` defaults to the same-named layout in the
    destination, else Blank. Image/chart relationship ids are remapped.

    Reliable for text, pictures and the live charts these decks carry; exotic OLE objects
    and notes slides are not guaranteed to transfer — a human reviews the compiled deck.
    """
    if layout is None:
        try:
            layout = _layout(dest_prs, src_slide.slide_layout.name)
        except Exception:
            layout = _layout(dest_prs, L_BLANK)
    dest = dest_prs.slides.add_slide(layout)
    for sh in list(dest.shapes):            # drop cloned empty placeholders
        _drop(sh)
    for shp in src_slide.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    rmap = {}
    for rId, rel in src_slide.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        try:
            if rel.is_external:
                rmap[rId] = dest.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            elif "image" in rel.reltype:
                # add by blob so the destination assigns a FRESH media partname
                # (relate_to keeps the source's partname and collides across decks)
                _img, rmap[rId] = dest.part.get_or_add_image_part(
                    io.BytesIO(rel.target_part.blob))
            else:
                rmap[rId] = dest.part.relate_to(rel.target_part, rel.reltype)
        except Exception:
            continue
    if rmap:
        for el in dest.shapes._spTree.iter():
            for attr in ("embed", "link", "id"):          # 'id' covers chart/OLE refs
                v = el.get(_R + attr)
                if v in rmap:
                    el.set(_R + attr, rmap[v])
    return dest


def has_native_chart(slide) -> bool:
    return any(getattr(sh, "has_chart", False) for sh in slide.shapes)


# ───────────────────────── fit to the safe area ─────────────────────────
_EMU_PER_PT = 12700
_FIT_MIN_VISUAL_H = Inches(1.0)


def _run_size_pt(paragraph, default=18.0):
    for r in paragraph.runs:
        if r.font.size is not None:
            return r.font.size.pt
    return default


def _text_overflows_box(shape) -> bool:
    """Rough estimate: does the rendered text need more height than the box provides?
    python-pptx cannot measure rendered extent, so approximate lines x line-height."""
    try:
        W, H = int(shape.width), int(shape.height)
    except Exception:
        return False
    if W <= 0 or H <= 0:
        return False
    total = 0
    for p in shape.text_frame.paragraphs:
        size = _run_size_pt(p)
        line_h = size * 1.55 * _EMU_PER_PT            # recipe body leading
        char_w = size * 0.52 * _EMU_PER_PT
        chars = max(1, len(p.text or ""))
        lines = max(1, -(-int(chars * char_w) // W))
        total += lines * line_h
    return total > H * 1.05


def _shrink_font(text_frame, factor=0.85, floor=12.0, default=18.0):
    """Shrink a text frame's runs. Floor is 12pt = the recipe's 16px source-line size —
    body text should never go below it; split the slide instead."""
    for p in text_frame.paragraphs:
        for r in p.runs:
            cur = r.font.size.pt if r.font.size is not None else default
            r.font.size = Pt(max(floor, round(cur * factor, 1)))


def fit_above_footer(slide) -> int:
    """Nudge/shrink shapes so nothing crosses the template's footer band (6.82 in).
    Captions and short labels move up; pictures/charts/tables shrink from the bottom;
    dense text boxes whose text overflows get a smaller font. Nothing is moved above the
    title zone. Returns the number of shapes adjusted."""
    safe, tz = int(FOOTER_SAFE_BOTTOM), int(TITLE_ZONE_TOP)
    changed = 0
    for sh in slide.shapes:
        try:
            top, h = int(sh.top), int(sh.height)
        except Exception:
            continue
        if h <= 0:
            continue
        is_text = bool(getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip())
        is_visual = (sh.shape_type == MSO_SHAPE_TYPE.PICTURE
                     or getattr(sh, "has_chart", False) or getattr(sh, "has_table", False))
        overflow = (top + h) - safe
        if overflow > 0:
            if is_visual:
                sh.height = max(int(_FIT_MIN_VISUAL_H), h - overflow)
                changed += 1
            elif is_text and h < int(Inches(1.0)):
                sh.top = max(tz, top - overflow)
                changed += 1
            elif is_text:
                sh.top = max(tz, top - overflow)
                if int(sh.top) + h > safe:
                    _shrink_font(sh.text_frame)
                changed += 1
        elif is_text and _text_overflows_box(sh):
            _shrink_font(sh.text_frame)
            changed += 1
    return changed


# ───────────────────────── output ─────────────────────────
def save_deck(prs: Presentation, name: str, out_dir: Path | None = None) -> Path:
    out_dir = Path(out_dir) if out_dir else default_output_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    if not name.lower().endswith(".pptx"):
        name += ".pptx"
    path = out_dir / name
    prs.save(str(path))
    return path


# ───────────────────── figures: the figs/ pipeline ─────────────────────
# Charts arrive as images and this skill does not care what drew them — matplotlib
# (gl-matplotlib), ggplot (gl-ggplot's `save_fig`/`gl_export_fig`), or anything else.
FIG_SIZES = {                                  # inches; mirror gl-ggplot's named sizes
    "slide": (10, 5.625),                      # 16:9 — the whole slide's aspect
    "slide_half": (4.9, 5.0),                  # side-by-side pair on a cols slide
    "slide_wide": (11.5, 4.4),                 # wide/short, e.g. a ranked bar chart
    # Fills the Single Visual chart area (12.4 x 5.42 in = aspect 2.29, wider than 16:9),
    # so a full-slide chart has no dead margin either side. Prefer this on a chart slide.
    "slide_fill": (12.4, 5.4),
    "full": (6.5, 4.0), "full_tall": (6.5, 6.0), "full_square": (6.5, 6.5),
    "half": (3.167, 3.0),
}
DEFAULT_LANG = os.environ.get("GL_FIG_LANG", "en")


def figs_root(explicit: Path | None = None) -> Path:
    """Where exported figures live: explicit arg -> $GL_FIGS_DIR -> the nearest `figs/`
    directory at or above the cwd -> ./figs. No project layout is assumed."""
    if explicit:
        return Path(explicit)
    env = os.environ.get("GL_FIGS_DIR")
    if env:
        return Path(env)
    here = Path.cwd()
    for cand in [here, *here.parents]:
        if (cand / "figs").is_dir():
            return cand / "figs"
    return here / "figs"


def _figs_dir(group: str, root: Path | None = None) -> Path:
    return figs_root(root) / group


def render_figure(fig, out_png: Path | None = None, name: str = "figure",
                  dpi: int = 200) -> Path:
    """Save a matplotlib figure as a slide-ready PNG (ad hoc — no manifest entry)."""
    out = Path(out_png) if out_png else (default_output_dir() / "_assets" / (name + ".png"))
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(str(out), dpi=dpi, bbox_inches="tight", facecolor=PAPER["paper"])
    return out


def export_fig(fig, group: str, name: str, *, title: str = None, source: str = None,
               caption: str = None, lang: str = None, size="slide", dpi: int = 200,
               root: Path | None = None) -> Path:
    """Save a figure as a slide-ready PNG and record it in `figures.json`.

    Writes `<figs>/<group>/<name>_<lang>.png`. `group` is whatever bucket suits the
    project (a script name, a notebook name, a chapter). The call belongs in the cell or
    function that BUILDS the chart, so figures regenerate when the analysis re-runs.

    Language variants are first-class: export the same `name` once per language and the
    deck builder picks the one it needs via `find_fig(..., lang=...)`.
    """
    if size is not None:
        wh = FIG_SIZES[size] if isinstance(size, str) else size
        fig.set_size_inches(*wh)
    lang = lang or DEFAULT_LANG
    stem = name + "_" + lang if lang else name
    d = _figs_dir(group, root)
    d.mkdir(parents=True, exist_ok=True)
    out = d / (stem + ".png")
    fig.savefig(str(out), dpi=dpi, bbox_inches="tight", facecolor=PAPER["paper"])
    manifest = d / "figures.json"
    data = {}
    if manifest.exists():
        try:
            data = json.loads(manifest.read_text(encoding="utf-8"))
        except Exception:
            data = {}
    data[stem] = {"name": name, "file": out.name, "lang": lang, "title": title,
                  "source": source, "caption": caption}
    manifest.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return out


def find_fig(group: str, name: str = None, lang: str = None, root: Path | None = None):
    """Look up exported figure(s). With `name`, the best match as a dict (path, title,
    source, caption, lang); without it, every entry in the group. Falls back to globbing
    the directory when no manifest is present."""
    d = _figs_dir(group, root)
    if not d.exists():
        return None if name else []
    manifest = d / "figures.json"
    entries = {}
    if manifest.exists():
        try:
            entries = json.loads(manifest.read_text(encoding="utf-8"))
        except Exception:
            entries = {}
    if not entries:
        for png in sorted(d.glob("*.png")):
            stem = png.stem
            lg = stem.rsplit("_", 1)[1] if "_" in stem and len(stem.rsplit("_", 1)[1]) == 2 else None
            base = stem[: -(len(lg) + 1)] if lg else stem
            entries[stem] = {"name": base, "file": png.name, "lang": lg,
                             "title": None, "source": None, "caption": None}

    def _with_path(stem, e):
        return dict(e, path=str(d / e["file"]), stem=stem)

    if name is None:
        return [_with_path(s, e) for s, e in entries.items()]
    cands = [(s, e) for s, e in entries.items() if e.get("name") == name]
    if not cands:
        return None
    want = lang or DEFAULT_LANG
    for s, e in cands:
        if e.get("lang") == want:
            return _with_path(s, e)
    return _with_path(*cands[0])


# ───────────────────────── validation ─────────────────────────
def _allowed_colors() -> set:
    allowed = {"FFFFFF", "000000"}
    for group in (INK, ACCENT, PAPER, MUTED):
        for v in group.values():
            allowed.add(v.lstrip("#").upper())
    for tones in CAT.values():
        for v in tones.values():
            allowed.add(v.lstrip("#").upper())
    for ramp in list(SEQUENTIAL.values()) + list(DIVERGING.values()):
        for v in ramp:
            allowed.add(v.lstrip("#").upper())
    return allowed


ALLOWED_FONTS = {None, "Source Sans Pro", "Source Sans 3", "+mj-lt", "+mn-lt"}

# Layouts with a dark ground (#124560 in the shipped template). Their placeholder text
# inherits white from the master, so builders must not set a dark color on them.
DARK_LAYOUTS = {L_TITLE.strip().lower(), L_TITLE_BG.strip().lower(),
                L_CLOSING.strip().lower()}


def _luminance(rgb_hex: str) -> float:
    """Relative luminance, 0 (black) to 1 (white) — used only for the dark-on-dark check."""
    h = str(rgb_hex).lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def validate_deck(path, strict: bool = False, quiet: bool = False) -> dict:
    """Re-open a finished deck and check it is valid, on-template and on-grammar:

      - the file opens and is 16:9;
      - every slide sits on one of the template's 12 layouts;
      - no run declares an off-template font (builders should declare none at all);
      - every explicit text color resolves to a grammar token;
      - nothing crosses the footer band;
      - no empty placeholder is left showing its prompt;
      - chart-slide titles end with a period (recipe rule 5) — reported as a NOTE.
    """
    report = {"path": str(path), "ok": True, "n_slides": 0, "issues": [], "notes": []}
    try:
        prs = Presentation(str(path))
    except Exception as e:
        report["ok"] = False
        report["issues"].append("FILE: cannot open deck (" + str(e) + ")")
        if not quiet:
            print(_format_report(report))
        if strict:
            raise
        return report

    report["n_slides"] = len(prs.slides)
    if abs(int(prs.slide_width) - int(SLIDE_W)) > 50000 or \
       abs(int(prs.slide_height) - int(SLIDE_H)) > 50000:
        report["issues"].append("SIZE: slide is not 16:9 ("
                                + str(prs.slide_width) + "x" + str(prs.slide_height) + " EMU)")

    known = {n.strip().lower() for n in TEMPLATE_LAYOUTS}
    allowed_colors = _allowed_colors()
    safe = int(FOOTER_SAFE_BOTTOM)
    for n, slide in enumerate(prs.slides, start=1):
        sn = "s" + str(n)
        lay = slide.slide_layout.name
        if lay.strip().lower() not in known:
            report["issues"].append(
                "LAYOUT " + sn + ": '" + lay + "' is not one of the template's 12 layouts")
        if lay.strip() == L_SINGLE.strip():
            for ph in slide.placeholders:
                if ph.placeholder_format.idx == 14 and ph.text_frame.text.strip():
                    t = ph.text_frame.text.strip()
                    if not t.endswith((".", "?", "!")):
                        report["notes"].append(
                            "TITLE " + sn + ": chart titles end with a period (recipe rule 5): "
                            + repr(t[:60]))
        for shape in slide.shapes:
            try:
                bottom = int(shape.top) + int(shape.height)
            except Exception:
                bottom = 0
            if bottom > safe + 5000:
                report["issues"].append(
                    "FIT " + sn + ": a " + str(shape.shape_type)
                    + " crosses the footer band (bottom "
                    + "{:.2f}".format(bottom / 914400) + " in > 6.82 in)")
            if shape.is_placeholder and getattr(shape, "has_text_frame", False) \
                    and not shape.text_frame.text.strip() \
                    and shape.placeholder_format.type in _PIC_TYPES:
                report["issues"].append(
                    "EMPTY " + sn + ": an unfilled picture placeholder will show its prompt")
            if not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    fn = run.font.name
                    if fn is not None and fn not in ALLOWED_FONTS:
                        report["issues"].append(
                            "FONT " + sn + ": off-template font '" + str(fn) + "' in "
                            + repr(run.text[:30]))
                    try:
                        rgb = run.font.color.rgb
                    except Exception:
                        rgb = None
                    if rgb is not None and str(rgb).upper() not in allowed_colors:
                        report["issues"].append(
                            "COLOR " + sn + ": off-grammar #" + str(rgb) + " in "
                            + repr(run.text[:30]))
                    if rgb is not None and lay.strip().lower() in DARK_LAYOUTS \
                            and _luminance(str(rgb)) < 0.4:
                        report["issues"].append(
                            "CONTRAST " + sn + ": dark text (#" + str(rgb) + ") on the "
                            "dark '" + lay.strip() + "' layout - use paper, or let it "
                            "inherit: " + repr(run.text[:30]))
    report["ok"] = not report["issues"]
    if not quiet:
        print(_format_report(report))
    if strict and not report["ok"]:
        raise ValueError("Deck failed GL validation: "
                         + str(len(report["issues"])) + " issue(s).")
    return report


def _format_report(report: dict) -> str:
    head = "OK - valid, on-template, on-grammar" if report["ok"] \
        else "FAIL - " + str(len(report["issues"])) + " issue(s)"
    lines = ["[validate_deck] " + report["path"],
             "  slides: " + str(report["n_slides"]) + "   " + head]
    for issue in report["issues"][:40]:
        lines.append("    - " + issue)
    if len(report["issues"]) > 40:
        lines.append("    ... and " + str(len(report["issues"]) - 40) + " more")
    for note in report.get("notes", [])[:20]:
        lines.append("    ~ " + note)
    return "\n".join(lines)


if __name__ == "__main__":
    drift = check_token_drift()
    print("token drift:", "OK (" + str(drift["checked"]) + " hexes match grammar.md)"
          if drift["ok"] else drift["missing"])
    prs = new_deck()
    print("layouts:", layout_names(prs))
    add_break_slide(prs, "A section divider")
    add_content_slide(prs, "A content slide",
                      ["First point", "Second point", ("A nested detail", 1)],
                      source="Source: example.")
    import tempfile
    out = save_deck(prs, "_smoke_test", out_dir=Path(tempfile.gettempdir()))
    print("wrote", out)
    validate_deck(out)
